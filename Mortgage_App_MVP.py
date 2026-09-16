import pandas as pd
import numpy as np
from datetime import datetime
from dateutil.relativedelta import relativedelta
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.dataframe import dataframe_to_rows
import tempfile
import os
import streamlit as st


def calculate_monthly_payment(principal, annual_rate, term_months):
    monthly_rate = annual_rate / 12
    payment = principal * (monthly_rate * (1 + monthly_rate) ** term_months) / \
              ((1 + monthly_rate) ** term_months - 1)
    return round(payment, 2)


def generate_amortization_schedule(principal, annual_rate, years, start_date, prepayments=None):
    monthly_rate = annual_rate / 12
    term_months = years * 12
    payment = calculate_monthly_payment(principal, annual_rate, term_months)
    schedule = []

    current_balance = principal
    current_date = pd.to_datetime(start_date)
    new_payment = payment

    if prepayments is None:
        prepayments = {}

    for month in range(1, term_months + 1):
        interest_payment = round(current_balance * monthly_rate, 2)
        capital_payment = round(new_payment - interest_payment, 2)
        prepayment = prepayments.get(current_date.strftime("%Y-%m"), 0)

        # Ensure final payment doesn't overpay
        if current_balance - capital_payment - prepayment < 0:
            capital_payment = current_balance - prepayment
            new_payment = capital_payment + interest_payment

        current_balance = round(current_balance - capital_payment - prepayment, 2)

        schedule.append({
            'Payment Date': current_date.strftime("%Y-%m-%d"),
            'Original Installment': payment,
            'Interest Payment': interest_payment,
            'Capital Payment': capital_payment,
            'Prepayment': prepayment,
            'New Installment': new_payment,
            'Installment Difference': round(payment - new_payment, 2),
            'Remaining Balance': current_balance
        })

        # Recalculate new installment if prepayment is made
        if prepayment > 0 and current_balance > 0:
            remaining_months = term_months - month
            new_payment = calculate_monthly_payment(current_balance, annual_rate, remaining_months)

        current_date += relativedelta(months=1)

        if current_balance <= 0:
            break

    df = pd.DataFrame(schedule)
    return df


def get_user_inputs():
    print("=== Mortgage Calculator ===")

    while True:
        try:
            principal = float(input("Enter the mortgage amount (e.g. 300000): "))
            if principal <= 0:
                raise ValueError
            break
        except ValueError:
            print("Please enter a valid positive number.")

    while True:
        try:
            annual_rate = float(input("Enter the annual interest rate (as % e.g. 4 for 4%): "))
            if annual_rate <= 0 or annual_rate > 100:
                raise ValueError
            annual_rate = annual_rate / 100  # convert to decimal
            break
        except ValueError:
            print("Please enter a valid percentage between 0 and 100.")

    while True:
        try:
            years = int(input("Enter the mortgage term in years (e.g. 30): "))
            if years <= 0:
                raise ValueError
            break
        except ValueError:
            print("Please enter a valid number of years.")

    while True:
        try:
            start_date = input("Enter the start date (YYYY-MM-DD): ")
            datetime.strptime(start_date, "%Y-%m-%d")
            break
        except ValueError:
            print("Please enter a valid date in YYYY-MM-DD format.")

    # Collect prepayments
    prepayments = {}
    add_prepayments = input("Do you want to add prepayments? (yes/no): ").strip().lower()

    while add_prepayments == 'yes':
        try:
            prepay_date = input("Enter prepayment date (YYYY-MM): ").strip()
            datetime.strptime(prepay_date, "%Y-%m")  # validate format
            amount = float(input(f"Enter prepayment amount for {prepay_date}: "))
            if amount <= 0:
                raise ValueError
            prepayments[prepay_date] = prepayments.get(prepay_date, 0) + amount
        except ValueError:
            print("Invalid date or amount. Please try again.")
            continue

        add_more = input("Add another prepayment? (yes/no): ").strip().lower()
        if add_more != 'yes':
            break

    return principal, annual_rate, years, start_date, prepayments



#########################################################################################################
#Extract into excel
def export_to_excel(principal, annual_rate, years, start_date, prepayments, schedule_df, filename="Mortgage_Schedule.xlsx"):
    # Create a new Excel workbook
    wb = Workbook()
    del wb['Sheet']  # remove default sheet

    ### Tab 1: Loan Contract ###
    ws1 = wb.create_sheet("Loan Contract")
    ws1.append(["Parameter", "Value"])
    ws1.append(["Mortgage Amount", principal])
    ws1.append(["Annual Interest Rate (%)", annual_rate * 100])
    ws1.append(["Term (years)", years])
    ws1.append(["Start Date", start_date])
    ws1.append(["Monthly Payment (Original)", calculate_monthly_payment(principal, annual_rate, years * 12)])

    ### Tab 2: Prepayments ###
    ws2 = wb.create_sheet("Prepayments")
    ws2.append(["Date (YYYY-MM)", "Amount"])
    for date, amount in sorted(prepayments.items()):
        ws2.append([date, amount])

    ### Tab 3: Amortization Schedule ###
    ws3 = wb.create_sheet("Amortization Schedule")
    for r in dataframe_to_rows(schedule_df, index=False, header=True):
        ws3.append(r)

    ### Tab 4: Graphs ###
    # Create a plot comparing original vs new installment amounts
    fig, ax = plt.subplots(figsize=(10, 5))
    dates = pd.to_datetime(schedule_df["Payment Date"])
    original = schedule_df["Original Installment"]
    new = schedule_df["New Installment"]

    ax.plot(dates, original, label="Original Installment", color='blue')
    ax.plot(dates, new, label="New Installment (with prepayments)", color='green')
    ax.set_title("Monthly Installments: Original vs With Prepayments")
    ax.set_xlabel("Date")
    ax.set_ylabel("Installment Amount")
    ax.legend()
    ax.grid(True)
    fig.tight_layout()

    # Save chart temporarily and insert into Excel
    tmp_img = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    chart_path = tmp_img.name
    plt.savefig(chart_path)
    plt.close()

    ws4 = wb.create_sheet("Graphs")
    img = XLImage(chart_path)
    img.anchor = "A1"
    ws4.add_image(img)

    # Save workbook
    wb.save(filename)
    print(f"\n✅ Excel export complete: {filename}")

    # Clean up temporary image
    os.remove(chart_path)


#########################################################################################################
#MAIN CALL

if __name__ == "__main__":
    principal, annual_rate, years, start_date, prepayments = get_user_inputs()

    schedule_df = generate_amortization_schedule(
        principal, annual_rate, years, start_date, prepayments
    )

    pd.set_option('display.max_rows', None)
    print("\n=== Mortgage Payment Schedule ===\n")
    print(schedule_df)

    # Export to Excel
    export_to_excel(principal, annual_rate, years, start_date, prepayments, schedule_df)

#########################################################################################################

# Example usage
#if __name__ == "__main__":
#    # Inputs
#    principal = 300000        # Mortgage amount
#    annual_rate = 0.04        # 4% annual interest
#    years = 30                # 30-year mortgage
#    start_date = '2025-09-01' # Start of mortgage

#    # Optional: Add prepayments here (format: 'YYYY-MM': amount)
#    prepayments = {
#       '2026-09': 10000,
#        '2028-01': 15000
#    }

#    schedule_df = generate_amortization_schedule(principal, annual_rate, years, start_date, prepayments)

# Display result
    #pd.set_option('display.max_rows', None)  # Show full table if needed
    #print("\n=== Mortgage Payment Schedule ===\n")
    #print(schedule_df)
